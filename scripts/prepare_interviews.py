#!/usr/bin/env python3
"""Stage 7 transcript preparation and shared qualitative schema creation."""

from __future__ import annotations

import collections
import re
from pathlib import Path

import pandas as pd
from docx import Document

from pipeline_common import INTERVIEW_DIR, ROOT, write_csv

PROCESSED = ROOT / "data/processed"
REPORTS = ROOT / "outputs/reports"


QUALITATIVE_CODES = [
    ("MIGRATION_CONTEXT", "迁移背景", "", "迁居、插班节点及其制度位置", "材料描述迁居或插班背景", "只描述一般学校生活", "DEDUCTIVE"),
    ("PRIOR_RESOURCES", "既有资源", "", "迁居前语言、学习、关系或家庭资源", "既有经验被参与者用于适应", "迁居后新获得的支持", "DEDUCTIVE"),
    ("SYSTEM_LANGUAGE_DISCONTINUITY", "教学语言断裂", "教育制度断裂", "粤语、英语、繁体字或学科语言造成的制度接口变化", "语言变化影响课堂理解或参与", "纯粹同伴幽默边界", "DEDUCTIVE"),
    ("SYSTEM_ASSESSMENT_RULES", "评价规则变化", "教育制度断裂", "作业、考试、汇报或评分要求改变", "参与者比较或学习评价要求", "没有规则内容的一般学业压力", "DEDUCTIVE"),
    ("HIDDEN_CURRICULUM", "隐性规则", "教育制度断裂", "未明确说明但需逐渐识别的学校规则", "参与者不知道默认程序或标准", "已明确告知的正式规定", "DEDUCTIVE"),
    ("PEER_LANGUAGE_HUMOUR_BOUNDARY", "语言与幽默边界", "同伴文化边界", "语言、口音、俗语或幽默与差别化参与相连", "具体互动显示难以进入谈话或群体", "仅报告语言能力差异而无参与后果", "DEDUCTIVE"),
    ("PEER_LOCAL_KNOWLEDGE_BOUNDARY", "地方知识边界", "同伴文化边界", "地方经验、学校传统或共同记忆影响参与", "地方知识与插不上话等后果相连", "一般兴趣不同", "DEDUCTIVE"),
    ("LABELING_ASSUMPTION", "标签与预设", "同伴文化边界", "来源、口音或内地经历触发预设或差别对待", "参与者明确描述预设、刻板印象或分组", "研究者推断他人态度", "DEDUCTIVE"),
    ("PARTICIPATION_WITHDRAWAL", "减少参与", "适应策略与能动性", "担忧、陌生或边界使参与者减少主动互动", "参与者明确描述沉默、回避或退出", "单纯缺少机会", "DEDUCTIVE"),
    ("PEER_INVITATION", "同伴邀请", "适应策略与能动性", "同伴主动邀请进入互动或活动", "有具体邀请行动", "泛称同伴友好", "DEDUCTIVE"),
    ("GROUPWORK_ENTRY", "小组活动进入关系", "适应策略与能动性", "小组作业、汇报或活动成为关系进入点", "具体小组事件改变互动", "只描述课程形式", "DEDUCTIVE"),
    ("MAINLAND_PEER_SUPPORT", "内地背景同伴支持", "支持资源", "具有内地经历的同伴提供信息、情绪或关系支持", "支持者背景及行动均明确", "研究者按语言猜测背景", "DEDUCTIVE"),
    ("HK_PEER_SUPPORT", "香港本地同伴支持", "支持资源", "本地同伴提供关系、学业或信息支持", "有明确支持行动", "仅共同在场", "DEDUCTIVE"),
    ("TEACHER_SUPPORT", "教师支持", "支持资源", "教师解释、安排、关怀或提供机会", "具体教师行动及作用", "一般学校制度", "DEDUCTIVE"),
    ("FAMILY_SUPPORT", "家庭支持", "支持资源", "家庭成员提供情绪、信息或物质支持", "材料明确家庭行动", "只说明迁居由家庭决定", "DEDUCTIVE"),
    ("SELF_INITIATED_ADAPTATION", "主动适应", "适应策略与能动性", "参与者主动练习、接近、询问或调整策略", "具体自主行动", "只有结果没有行动", "DEDUCTIVE"),
    ("RULE_LEARNING", "规则学习", "适应策略与能动性", "通过观察、解释、试错或支持掌握学校规则", "材料呈现学习规则的过程", "静态规则理解评分", "DEDUCTIVE"),
    ("LEARNER_IDENTITY_RENEGOTIATION", "学习者身份重构", "归属与身份结果", "参与者重新理解能力、优势或好学生标准", "明确区分规则陌生与能力或描述标准变化", "研究者推断自尊", "DEDUCTIVE"),
    ("BELONGING", "被接纳与归属", "归属与身份结果", "参与者感到成为班级成员、可自然参与或有人信任", "具体感受或互动后果", "只报告平台加入", "DEDUCTIVE"),
    ("EXCLUSION", "排斥或孤立", "归属与身份结果", "参与者明确描述被排除、孤立或差别参与", "存在事件或直接报告", "未被提及", "DEDUCTIVE"),
    ("BROKERAGE_ACTION", "解释翻译与连接行动", "文化中介", "参与者在不同背景者之间解释、翻译或连接", "有具体行为", "仅位于多群体网络", "DEDUCTIVE"),
    ("BROKERAGE_BENEFIT", "中介收益", "文化中介", "中介行动带来认可、关系或能力收益", "行动与收益均有证据", "只发生一次解释", "DEDUCTIVE"),
    ("BROKERAGE_PRESSURE", "中介负担", "文化中介", "解释者或中间人角色带来疲惫、压力或边缘感", "参与者明确报告负担", "研究者推断", "DEDUCTIVE"),
    ("PLATFORM_OLD_TIES", "平台维持旧关系", "数字平台实践", "平台用于维持迁居前关系或获得支持", "关系、平台和功能明确", "只列平台名称", "DEDUCTIVE"),
    ("PLATFORM_HK_SCHOOL", "平台参与香港学校关系", "数字平台实践", "平台用于班群、学校信息或本地同伴互动", "平台功能与学校关系明确", "只安装平台", "DEDUCTIVE"),
    ("PLATFORM_DIFFERENTIATION", "平台功能分工", "数字平台实践", "多个平台按关系、受众或功能并存", "至少两类关系或用途的对照", "单平台使用", "DEDUCTIVE"),
    ("PLATFORM_RULE_BARRIER", "线上规则门槛", "数字平台实践", "不熟悉平台或班群规则造成信息或参与障碍", "存在具体线上后果", "一般技术困难", "DEDUCTIVE"),
    ("TURNING_POINT", "转折点", "", "改变适应、关系或规则理解的关键事件", "材料呈现事件前后变化", "研究者仅按时间切段", "DEDUCTIVE"),
    ("SURVEY_INTERVIEW_DIVERGENCE", "问卷访谈分歧", "矛盾与反例", "两类资料对同一主题指向不同", "可定位具体问卷项与访谈段", "时点不同但方向不冲突", "DEDUCTIVE"),
    ("NON_CULTURAL_ALTERNATIVE", "非文化替代解释", "矛盾与反例", "学校类型、教学语言、年级、个性、家庭资源或班级结构等解释", "材料支持或合理限定的替代机制", "无证据猜测心理状态", "DEDUCTIVE"),
]


RESTRICTED_REDACTION_TERMS_PATH = PROCESSED / "restricted_redaction_terms.txt"


def restricted_location_terms() -> list[str]:
    if not RESTRICTED_REDACTION_TERMS_PATH.is_file():
        raise FileNotFoundError(
            "Restricted redaction-term file is required locally and must remain outside Git"
        )
    return [
        term.strip()
        for term in RESTRICTED_REDACTION_TERMS_PATH.read_text(encoding="utf-8").splitlines()
        if term.strip() and not term.lstrip().startswith("#")
    ]


def redact_restricted_locations(text: str) -> tuple[str, bool]:
    terms = {term for term in restricted_location_terms() if term in text}
    for term in sorted(terms, key=len, reverse=True):
        text = text.replace(term, "[内地粤语地区]")
    return text, bool(terms)


def redact_text(text: str, sensitive_labels: set[str] | None = None) -> tuple[str, list[str]]:
    changes: list[str] = []
    new_text = re.sub(
        r"^\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2}\s*",
        "",
        text,
    )
    if new_text != text:
        text = new_text
        changes.append("exact_timestamp")
    for label in sensitive_labels or set():
        if label in text:
            text = text.replace(label, "[说话人标识已隐去]")
            changes.append("speaker_name_or_handle")
    new_text = re.sub(
        r"(?:我的年龄是|我年龄是|年龄是)?\s*(?:嗯\s*)?(?:1[89]|[2-9]\d)\s*岁",
        "[具体年龄已泛化：成年]",
        text,
    )
    if new_text != text:
        text = new_text
        changes.append("exact_age")
    new_text = re.sub(r"\b1\d{10}\b", "[电话号码已隐去]", text)
    if new_text != text:
        text = new_text
        changes.append("phone")
    new_text = re.sub(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        "[电子邮箱已隐去]",
        text,
    )
    if new_text != text:
        text = new_text
        changes.append("email")
    text, location_changed = redact_restricted_locations(text)
    if location_changed:
        changes.append("specific_location")
    return text.strip(), changes


def infer_role_maps() -> dict[Path, dict[str, str]]:
    label_pattern = re.compile(r"^(.*?)\s+\d{4}-\d{2}-\d{2}\s+\d\d:\d\d:\d\d$")
    counts_by_path: dict[Path, collections.Counter[str]] = {}
    files_by_label: collections.Counter[str] = collections.Counter()
    for path in sorted(INTERVIEW_DIR.glob("*.docx")):
        counts: collections.Counter[str] = collections.Counter()
        for paragraph in Document(path).paragraphs:
            match = label_pattern.match(paragraph.text.strip())
            if match:
                counts[match.group(1)] += 1
        if counts:
            counts_by_path[path] = counts
            files_by_label.update(counts.keys())
    shared_interviewers = {label for label, count in files_by_label.items() if count > 1}
    role_maps: dict[Path, dict[str, str]] = {}
    for path, counts in counts_by_path.items():
        mapping = {label: "Interviewer" for label in counts if label in shared_interviewers}
        remaining = [label for label in counts if label not in mapping]
        if remaining:
            participant = max(remaining, key=lambda label: counts[label])
            mapping[participant] = "Participant"
            interviewer_number = 2
            for label in remaining:
                if label != participant:
                    mapping[label] = f"Interviewer {interviewer_number}"
                    interviewer_number += 1
        role_maps[path] = mapping
    return role_maps


def timestamped_lines(path: Path, role_map: dict[str, str]) -> tuple[list[tuple[str, str]], list[str]]:
    document = Document(path)
    label_pattern = re.compile(r"^(.*?)\s+\d{4}-\d{2}-\d{2}\s+\d\d:\d\d:\d\d$")
    current_role = "System"
    lines: list[tuple[str, str]] = []
    redactions: list[str] = []
    seen_speaker_labels: set[str] = set()
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        match = label_pattern.match(text)
        if match:
            raw_label = match.group(1)
            current_role = role_map.get(raw_label, "Unknown speaker")
            if raw_label not in seen_speaker_labels:
                redactions.append("speaker_name_or_handle")
                seen_speaker_labels.add(raw_label)
            continue
        if "会议已开启实时转写" in text:
            current_role = "System"
        cleaned, changes = redact_text(text, set(role_map))
        redactions.extend(changes)
        lines.append((current_role, cleaned))
    return lines, redactions


def summary_lines(path: Path, participant_id: str) -> tuple[list[tuple[str, str]], list[str]]:
    document = Document(path)
    lines: list[tuple[str, str]] = []
    redactions: list[str] = []
    current_role = "Source note"
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text.startswith("答："):
            current_role = "Participant"
            text = text[2:].strip()
        elif text.endswith(("？", "?")):
            current_role = "Interviewer"
        elif len(text) <= 18 and text.endswith(("：", ":")):
            current_role = "Source note"
        elif text.startswith("其他访谈补充信息"):
            current_role = "Source note"
        cleaned, changes = redact_text(text)
        redactions.extend(changes)
        lines.append((current_role, cleaned))
    return lines, redactions


def write_transcripts() -> None:
    mapping = pd.read_csv(PROCESSED / "interview_survey_map.csv")
    role_maps = infer_role_maps()
    report_rows = []
    for row in mapping.itertuples(index=False):
        source = INTERVIEW_DIR / row.interview_filename
        if row.participant_id == "P01":
            lines, redactions = summary_lines(source, row.participant_id)
            source_type = "整理式访谈记录；部分段落为来源笔记，不作逐字引语"
        else:
            lines, redactions = timestamped_lines(source, role_maps[source])
            source_type = "会议机器转写；已移除时间戳并匿名化说话人"
        output = PROCESSED / "interviews" / f"interview_{row.participant_id}_numbered.md"
        header = [
            f"# {row.participant_id} 编号访谈",
            "",
            f"- 来源类型：{source_type}",
            f"- 链接状态：{row.mapping_status}",
            "- 说明：保留原始措辞与机器识别错误；未进行语言润色。",
            "",
        ]
        body = [
            f"[L{line_number:04d}] {role}: {text}"
            for line_number, (role, text) in enumerate(lines, start=1)
        ]
        output.write_text("\n".join(header + body) + "\n", encoding="utf-8")
        report_rows.append(
            {
                "participant_id": row.participant_id,
                "source_file": f"RESTRICTED_SOURCE_{row.participant_id}.docx",
                "output_file": output.relative_to(ROOT).as_posix(),
                "numbered_lines": len(lines),
                "mapping_status": row.mapping_status,
                "redaction_events": len(redactions),
                "redaction_types": "|".join(sorted(set(redactions))) if redactions else "NONE",
                "source_type": source_type,
            }
        )
    write_csv(
        REPORTS / "transcript_preparation_details.csv",
        [
            "participant_id",
            "source_file",
            "output_file",
            "numbered_lines",
            "mapping_status",
            "redaction_events",
            "redaction_types",
            "source_type",
        ],
        report_rows,
    )
    matched_n = sum(r["mapping_status"] == "MATCHED" for r in report_rows)
    ineligible_n = sum(r["mapping_status"] == "INELIGIBLE" for r in report_rows)
    summary_n = sum(r["source_type"].startswith("整理式") for r in report_rows)
    machine_n = len(report_rows) - summary_n
    report = f"""# 访谈转录准备报告

## 处理结果

- 发现访谈文件：{len(report_rows)}
- 合资格链接访谈：{matched_n}
- 唯一链接但不合资格、仅保留受限审计副本：{ineligible_n}
- 整理式记录：{summary_n}
- 会议机器转写：{machine_n}

## 处理规则

- 移除会议转写中的精确时间戳和原始说话人姓名/平台昵称，统一为 `Interviewer`、`Interviewer 2`、`Participant`。
- 精确年龄泛化为“成年”；自动扫描电话号码和电子邮箱。
- 保留口语、重复、停顿词及明显机器识别错误，不静默改写参与者话语。
- 整理式记录只有标为 `Participant` 的行可作逐字引语，`Source note` 只可作为来源说明，不能放入引号。
- 机器实时转写在引用时必须保留“机器转写可能有误”的限定。
- 唯一链接但不合资格的访谈只保留受限脱敏审计副本，不进行实质编码。

## 不确定性与剩余风险

- 自动脱敏无法证明已经识别全部间接标识；最终报告只选取非稀有、研究必要的短引语。
- 机器转写未与录音核对，因此无法校正同音字、断句或说话人误分配。
- 本次没有接触或生成录音文件。
"""
    (REPORTS / "transcript_preparation_report.md").write_text(report, encoding="utf-8")


def write_qualitative_codebook() -> None:
    rows = [
        {
            "code_id": code_id,
            "code_name": name,
            "parent_code": parent,
            "definition": definition,
            "include_when": include,
            "exclude_when": exclude,
            "example": "",
            "source": source,
            "status": "ACTIVE",
            "notes": "Stage 7 shared deductive seed; case agents may suggest but not directly edit.",
        }
        for code_id, name, parent, definition, include, exclude, source in QUALITATIVE_CODES
    ]
    write_csv(
        PROCESSED / "qualitative_codebook.csv",
        [
            "code_id",
            "code_name",
            "parent_code",
            "definition",
            "include_when",
            "exclude_when",
            "example",
            "source",
            "status",
            "notes",
        ],
        rows,
    )


def main() -> None:
    (PROCESSED / "interviews").mkdir(parents=True, exist_ok=True)
    write_qualitative_codebook()
    write_transcripts()
    print("prepared qualitative codebook and numbered transcripts")


if __name__ == "__main__":
    main()
